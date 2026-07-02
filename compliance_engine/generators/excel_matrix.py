"""openpyxl compliance control matrix generator."""
import io
from datetime import datetime
from typing import Dict, Any, List

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, GradientFill
    )
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

NAVY_HEX  = "0F2D5E"
GOLD_HEX  = "C8A850"
GREEN_HEX = "10B981"
RED_HEX   = "EF4444"
AMBER_HEX = "F59E0B"
GRAY_HEX  = "6B7280"
BLUE_HEX  = "3B82F6"

STATUS_FILL = {
    "pass":         GREEN_HEX,
    "fail":         RED_HEX,
    "compensating": AMBER_HEX,
    "not_tested":   GRAY_HEX,
    "na":           BLUE_HEX,
}
RISK_FILL = {
    "high":   RED_HEX,
    "medium": AMBER_HEX,
    "low":    GREEN_HEX,
}


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _border():
    thin = Side(style="thin", color="D1D5DB")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def generate_control_matrix(
    audit: Dict[str, Any],
    controls: List[Dict],
    findings: List[Dict],
    org_name: str = "Kincaid IQ",
) -> bytes:
    if not OPENPYXL_OK:
        raise RuntimeError("openpyxl not installed")

    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.sheet_view.showGridLines = False
    ws_summary.column_dimensions["A"].width = 24
    ws_summary.column_dimensions["B"].width = 40

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
    label_font = Font(name="Calibri", bold=True, size=11, color=NAVY_HEX)
    value_font = Font(name="Calibri", size=11)

    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"] = f"{org_name} — Compliance Audit Report"
    ws_summary["A1"].font = Font(name="Calibri", bold=True, color="FFFFFF", size=16)
    ws_summary["A1"].fill = _fill(NAVY_HEX)
    ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 36

    ws_summary.merge_cells("A2:B2")
    ws_summary["A2"] = f"Framework: {audit.get('framework', '').upper()}  |  Generated: {datetime.utcnow().strftime('%B %d, %Y')}"
    ws_summary["A2"].font = Font(name="Calibri", color=GOLD_HEX, size=11)
    ws_summary["A2"].fill = _fill(NAVY_HEX)
    ws_summary["A2"].alignment = Alignment(horizontal="center")
    ws_summary.row_dimensions[2].height = 20

    rows = [
        ("Audit Name", audit.get("name", "N/A")),
        ("Auditor", audit.get("auditor", "N/A")),
        ("Status", audit.get("status", "N/A").upper()),
        ("Total Controls", audit.get("total_controls", 0)),
        ("Passed", audit.get("passed", 0)),
        ("Failed", audit.get("failed", 0)),
        ("Not Tested", audit.get("not_tested", 0)),
        ("Compliance Score", f"{audit.get('score_pct', 0):.1f}%"),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws_summary.row_dimensions[i].height = 18
        cell_l = ws_summary.cell(row=i, column=1, value=label)
        cell_v = ws_summary.cell(row=i, column=2, value=value)
        cell_l.font = label_font
        cell_v.font = value_font
        cell_l.fill = _fill("EFF6FF") if i % 2 == 0 else _fill("FFFFFF")
        cell_v.fill = _fill("EFF6FF") if i % 2 == 0 else _fill("FFFFFF")
        cell_l.border = _border()
        cell_v.border = _border()

    # ── Sheet 2: Controls ─────────────────────────────────────────────────────
    ws_ctrl = wb.create_sheet("Controls")
    ws_ctrl.sheet_view.showGridLines = False

    headers = ["Control ID", "Title", "Risk Level", "Result", "Notes", "Tested At"]
    col_widths = [14, 55, 12, 14, 40, 20]
    for i, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws_ctrl.cell(row=1, column=i, value=h)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        cell.fill = _fill(NAVY_HEX)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _border()
        ws_ctrl.column_dimensions[get_column_letter(i)].width = w
    ws_ctrl.row_dimensions[1].height = 22

    for row_num, c in enumerate(controls, start=2):
        result_val = c.get("result", "not_tested")
        risk_val = c.get("risk_level", "medium")
        row_data = [
            c.get("control_ref", ""),
            c.get("control_title", ""),
            risk_val.upper(),
            result_val.upper().replace("_", " "),
            c.get("notes", "") or "",
            str(c.get("tested_at", ""))[:16],
        ]
        fill_color = "FFFFFF" if row_num % 2 == 0 else "F9FAFB"
        for col_num, val in enumerate(row_data, start=1):
            cell = ws_ctrl.cell(row=row_num, column=col_num, value=val)
            cell.font = Font(name="Calibri", size=9)
            cell.fill = _fill(fill_color)
            cell.border = _border()
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_num == 3:  # Risk
                cell.fill = PatternFill("solid", fgColor=RISK_FILL.get(risk_val, GRAY_HEX))
                cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if col_num == 4:  # Result
                cell.fill = PatternFill("solid", fgColor=STATUS_FILL.get(result_val, GRAY_HEX))
                cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="top")
        ws_ctrl.row_dimensions[row_num].height = 18

    ws_ctrl.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws_ctrl.freeze_panes = "A2"

    # ── Sheet 3: Findings ─────────────────────────────────────────────────────
    if findings:
        ws_find = wb.create_sheet("Findings")
        ws_find.sheet_view.showGridLines = False
        f_headers = ["Severity", "Title", "Description", "Remediation", "Status", "Due Date", "Owner"]
        f_widths = [12, 40, 55, 55, 12, 14, 20]
        for i, (h, w) in enumerate(zip(f_headers, f_widths), start=1):
            cell = ws_find.cell(row=1, column=i, value=h)
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
            cell.fill = _fill(NAVY_HEX)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _border()
            ws_find.column_dimensions[get_column_letter(i)].width = w
        ws_find.row_dimensions[1].height = 22

        for row_num, f in enumerate(findings, start=2):
            sev = f.get("severity", "medium")
            sev_color = RISK_FILL.get(sev, GRAY_HEX)
            row_data = [
                sev.upper(),
                f.get("title", ""),
                f.get("description", ""),
                f.get("remediation", "") or "",
                f.get("status", "open").upper(),
                str(f.get("due_date", "")) or "",
                f.get("owner", "") or "",
            ]
            fill_color = "FFFFFF" if row_num % 2 == 0 else "FEF2F2"
            for col_num, val in enumerate(row_data, start=1):
                cell = ws_find.cell(row=row_num, column=col_num, value=val)
                cell.font = Font(name="Calibri", size=9)
                cell.fill = _fill(fill_color)
                cell.border = _border()
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if col_num == 1:
                    cell.fill = PatternFill("solid", fgColor=sev_color)
                    cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="top")
            ws_find.row_dimensions[row_num].height = 18
        ws_find.auto_filter.ref = f"A1:{get_column_letter(len(f_headers))}1"
        ws_find.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
